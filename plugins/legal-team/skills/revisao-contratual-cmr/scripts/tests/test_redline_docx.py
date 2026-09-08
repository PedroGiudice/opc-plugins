# -*- coding: utf-8 -*-
"""Testes da biblioteca de redline OOXML (marcas de revisão sobre .docx original).

Fixtures construídas com python-docx e XML injetado (marcas de terceiro,
numeração), para reproduzir o que as minutas reais trazem.
"""
import copy
import sys
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from redline_docx import Redline, texto_aceito, texto_rejeitado, verificar  # noqa: E402

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _add_runs(par, partes):
    for texto, negrito in partes:
        run = par.add_run(texto)
        run.bold = negrito or None


def _set_numpr(par, ilvl=0, numid=1):
    ppr = par._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl_el = OxmlElement("w:ilvl")
    ilvl_el.set(qn("w:val"), str(ilvl))
    numid_el = OxmlElement("w:numId")
    numid_el.set(qn("w:val"), str(numid))
    numpr.append(ilvl_el)
    numpr.append(numid_el)
    ppr.append(numpr)


def _inject_third_party_change(par, texto_inserido, texto_excluido):
    """Anexa ao parágrafo uma exclusão e uma inserção rastreadas de outro advogado."""
    xml_del = (
        '<w:del %s w:id="900" w:author="Alexandre Aguiar" w:date="2026-09-02T10:00:00Z">'
        '<w:r><w:delText xml:space="preserve">%s</w:delText></w:r></w:del>'
    ) % (nsdecls("w"), texto_excluido)
    xml_ins = (
        '<w:ins %s w:id="901" w:author="Alexandre Aguiar" w:date="2026-09-02T10:00:00Z">'
        '<w:r><w:t xml:space="preserve">%s</w:t></w:r></w:ins>'
    ) % (nsdecls("w"), texto_inserido)
    par._p.append(parse_xml(xml_del))
    par._p.append(parse_xml(xml_ins))


@pytest.fixture
def minuta(tmp_path):
    """Minuta com runs fragmentados, numeração, tabela e marca de terceiro."""
    doc = Document()
    p0 = doc.add_paragraph()
    _add_runs(p0, [("CLÁUSULA 1ª", True), (" – DO OBJETO", False)])
    p1 = doc.add_paragraph()
    _add_runs(p1, [("1.1. O registro será feito no prazo máximo de ", False), ("()", False), (" dias.", False)])
    _set_numpr(p1, ilvl=1)
    p2 = doc.add_paragraph()
    _add_runs(p2, [("1.2. A CONTRATADA poderá ceder este contrato a seu ", False), ("exclusivo critério", True), (".", False)])
    _set_numpr(p2, ilvl=1)
    p3 = doc.add_paragraph("2.1. Prazo de 15 dias para assinatura.")
    _inject_third_party_change(p3, " Prazo prorrogável uma vez.", " Sem prorrogação.")
    tabela = doc.add_table(rows=1, cols=2)
    tabela.rows[0].cells[0].text = "Assinatura A"
    tabela.rows[0].cells[1].text = "Assinatura B"
    doc.add_paragraph("3.1. Foro da Comarca de Belo Horizonte.")
    doc.core_properties.author = "Terceiro"
    caminho = tmp_path / "minuta.docx"
    doc.save(caminho)
    return caminho


def test_paragrafos_texto_vivo_une_runs_fragmentados(minuta):
    r = Redline(minuta, autor="CMR Advogados")
    textos = r.paragrafos()
    assert textos[0] == "CLÁUSULA 1ª – DO OBJETO"
    assert textos[1] == "1.1. O registro será feito no prazo máximo de () dias."
    # tabela é opaca: aparece como None, mas mantém a posição
    assert None in textos
    assert textos[-1] == "3.1. Foro da Comarca de Belo Horizonte."


def test_texto_vivo_ignora_marcas_de_terceiro(minuta):
    r = Redline(minuta, autor="CMR Advogados")
    # a inserção e a exclusão do terceiro ficam fora do texto vivo (não se edita dentro delas)
    assert r.paragrafos()[3] == "2.1. Prazo de 15 dias para assinatura."
    assert texto_aceito(minuta)[3] == "2.1. Prazo de 15 dias para assinatura. Prazo prorrogável uma vez."


def test_editar_gera_del_e_ins_com_autor(minuta, tmp_path):
    r = Redline(minuta, autor="CMR Advogados")
    r.editar("prazo máximo de", "()", "30 (trinta)")
    saida = tmp_path / "redline.docx"
    r.salvar(saida)
    est = r.estatisticas()
    assert est["ins"] == 2 and est["del"] == 2  # 1 par nosso + 1 par do terceiro (fixture)
    assert est["autores"]["CMR Advogados"] == 2
    assert texto_aceito(saida)[1] == "1.1. O registro será feito no prazo máximo de 30 (trinta) dias."
    assert texto_rejeitado(saida)[1] == "1.1. O registro será feito no prazo máximo de () dias."
    assert verificar(minuta, saida) == []


def test_editar_trecho_que_atravessa_runs(minuta, tmp_path):
    r = Redline(minuta, autor="CMR Advogados")
    r.editar("exclusivo critério", "a seu exclusivo critério.", "mediante anuência prévia e escrita da CONTRATANTE.")
    saida = tmp_path / "redline.docx"
    r.salvar(saida)
    assert texto_aceito(saida)[2] == "1.2. A CONTRATADA poderá ceder este contrato mediante anuência prévia e escrita da CONTRATANTE."
    assert texto_rejeitado(saida)[2] == "1.2. A CONTRATADA poderá ceder este contrato a seu exclusivo critério."
    assert verificar(minuta, saida) == []


def test_inserir_apos_ancora(minuta, tmp_path):
    r = Redline(minuta, autor="CMR Advogados")
    r.inserir("Foro da Comarca", "Belo Horizonte", ", Estado de Minas Gerais")
    saida = tmp_path / "redline.docx"
    r.salvar(saida)
    assert texto_aceito(saida)[-1] == "3.1. Foro da Comarca de Belo Horizonte, Estado de Minas Gerais."
    est = r.estatisticas()
    assert est["autores"]["CMR Advogados"] == 1  # só a inserção; o del que existe é do terceiro
    assert verificar(minuta, saida) == []


def test_substituir_paragrafo_inteiro(minuta, tmp_path):
    r = Redline(minuta, autor="CMR Advogados")
    r.substituir("Foro da Comarca", "3.1. Fica eleito o Foro da Comarca de Atibaia, Estado de São Paulo.")
    saida = tmp_path / "redline.docx"
    r.salvar(saida)
    assert texto_aceito(saida)[-1] == "3.1. Fica eleito o Foro da Comarca de Atibaia, Estado de São Paulo."
    assert texto_rejeitado(saida)[-1] == "3.1. Foro da Comarca de Belo Horizonte."


def test_excluir_paragrafo_marca_pilcrow_e_runs(minuta, tmp_path):
    r = Redline(minuta, autor="CMR Advogados")
    r.excluir("poderá ceder")
    saida = tmp_path / "redline.docx"
    r.salvar(saida)
    xml = zipfile.ZipFile(saida).read("word/document.xml").decode("utf-8")
    assert '<w:pPr>' in xml and '<w:rPr><w:del ' in xml  # marca de parágrafo excluída
    aceito = texto_aceito(saida)
    assert all("poderá ceder" not in (t or "") for t in aceito)
    assert texto_rejeitado(saida)[2] == "1.2. A CONTRATADA poderá ceder este contrato a seu exclusivo critério."


def test_novo_paragrafo_herda_numeracao_do_modelo(minuta, tmp_path):
    r = Redline(minuta, autor="CMR Advogados")
    r.novo_paragrafo(modelo="poderá ceder", texto="1.3. Nenhuma cessão dispensa a anuência da **CONTRATANTE**.")
    saida = tmp_path / "redline.docx"
    r.salvar(saida)
    aceito = texto_aceito(saida)
    assert aceito[3] == "1.3. Nenhuma cessão dispensa a anuência da CONTRATANTE."
    assert texto_rejeitado(saida)[3] is None or texto_rejeitado(saida)[3] == ""
    xml = zipfile.ZipFile(saida).read("word/document.xml").decode("utf-8")
    novo = xml.split("Nenhuma cessão")[0].rsplit("<w:p>", 1)[1]
    assert '<w:numPr><w:ilvl w:val="1"/>' in novo
    assert '<w:rPr><w:ins ' in novo  # marca de parágrafo inserida
    assert "<w:b/>" in xml.split("CONTRATANTE</w:t>")[0].rsplit("<w:r>", 1)[1]
    assert "**" not in "".join(t or "" for t in aceito)


def test_marcas_de_terceiro_permanecem_intactas(minuta, tmp_path):
    r = Redline(minuta, autor="CMR Advogados")
    antes = zipfile.ZipFile(minuta).read("word/document.xml").decode("utf-8")
    r.editar("prazo máximo de", "()", "30 (trinta)")
    saida = tmp_path / "redline.docx"
    r.salvar(saida)
    depois = zipfile.ZipFile(saida).read("word/document.xml").decode("utf-8")
    for trecho in ("Prazo prorrogável uma vez.", "Sem prorrogação."):
        bloco_antes = antes[antes.index(trecho) - 200: antes.index(trecho) + 40]
        assert trecho in depois
        assert 'w:author="Alexandre Aguiar"' in bloco_antes
    assert depois.count('w:author="Alexandre Aguiar"') == 2
    assert r.estatisticas()["autores"]["Alexandre Aguiar"] == 2


def test_salvar_preserva_pacote_e_ordem(minuta, tmp_path):
    r = Redline(minuta, autor="CMR Advogados")
    r.editar("prazo máximo de", "()", "30 (trinta)")
    saida = tmp_path / "redline.docx"
    r.salvar(saida)
    z1, z2 = zipfile.ZipFile(minuta), zipfile.ZipFile(saida)
    assert z1.namelist() == z2.namelist()
    for nome in z1.namelist():
        if nome != "word/document.xml":
            assert z1.read(nome) == z2.read(nome), nome


def test_localizar_exige_trecho_unico(minuta):
    r = Redline(minuta, autor="CMR Advogados")
    with pytest.raises(ValueError):
        r.localizar("CONTRATADA poderá ceder este contrato a seu")  # ok, único
        r.localizar("1.")  # ambíguo
    with pytest.raises(ValueError):
        r.localizar("texto que não existe")


def test_texto_novo_escapa_xml_e_preserva_espacos(minuta, tmp_path):
    r = Redline(minuta, autor="CMR Advogados")
    r.inserir("Foro da Comarca", "Belo Horizonte", " <sede da B&H> ")
    saida = tmp_path / "redline.docx"
    r.salvar(saida)
    xml = zipfile.ZipFile(saida).read("word/document.xml").decode("utf-8")
    assert "&lt;sede da B&amp;H&gt;" in xml
    assert 'xml:space="preserve"> &lt;sede' in xml
    assert texto_aceito(saida)[-1] == "3.1. Foro da Comarca de Belo Horizonte <sede da B&H> ."


def test_verificar_detecta_alteracao_sem_marca(minuta, tmp_path):
    adulterado = tmp_path / "adulterado.docx"
    with zipfile.ZipFile(minuta) as z_in, zipfile.ZipFile(adulterado, "w", zipfile.ZIP_DEFLATED) as z_out:
        for item in z_in.infolist():
            dados = z_in.read(item.filename)
            if item.filename == "word/document.xml":
                dados = dados.replace("Belo Horizonte".encode("utf-8"), "Atibaia".encode("utf-8"))
            z_out.writestr(item, dados)
    divergencias = verificar(minuta, adulterado)
    assert len(divergencias) == 1
    assert "Belo Horizonte" in divergencias[0]["original"]
    assert "Atibaia" in divergencias[0]["rejeitado"]


def test_marcador_de_pendencia_preservado(minuta, tmp_path):
    r = Redline(minuta, autor="CMR Advogados")
    r.editar("prazo máximo de", "()", "[● prazo] (● por extenso)")
    saida = tmp_path / "redline.docx"
    r.salvar(saida)
    assert "[● prazo] (● por extenso)" in texto_aceito(saida)[1]


def test_cli_aplica_operacoes_de_json(minuta, tmp_path):
    import json
    import subprocess

    ops = [
        {"op": "editar", "onde": "prazo máximo de", "de": "()", "para": "30 (trinta)"},
        {"op": "novo_paragrafo", "modelo": "poderá ceder", "texto": "1.3. Cláusula nova."},
        {"op": "excluir", "onde": "Foro da Comarca"},
    ]
    spec = tmp_path / "ops.json"
    spec.write_text(json.dumps(ops, ensure_ascii=False), encoding="utf-8")
    saida = tmp_path / "cli.docx"
    script = Path(__file__).resolve().parents[1] / "redline_docx.py"
    proc = subprocess.run(
        [sys.executable, str(script), "aplicar", str(minuta), str(spec), str(saida), "--autor", "Carlos Magno"],
        capture_output=True, text=True, check=False,
    )
    assert proc.returncode == 0, proc.stderr
    assert "ins=" in proc.stdout and "del=" in proc.stdout
    aceito = texto_aceito(saida)
    assert aceito[1].endswith("30 (trinta) dias.")
    assert aceito[3] == "1.3. Cláusula nova."
    assert verificar(minuta, saida) == []
    est = Redline(saida, autor="x").estatisticas()
    assert est["autores"]["Carlos Magno"] >= 4
